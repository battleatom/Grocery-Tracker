"""Convert retailer XLSX exports to normalized observation JSON.

Requires Python package openpyxl in CI/local import environment.
Usage: python scripts/import_xlsx.py input.xlsx output.json
"""
from __future__ import annotations
import json, sys
from pathlib import Path
from openpyxl import load_workbook
from urllib.parse import urlparse, parse_qs, unquote
import import_walmart, import_sams, import_smiths, import_albertsons, import_safeway

def retailer_for(row):
    raw=str(row.get("w-100 href") or row.get("citrus-Link href") or row.get("href") or row.get("url") or "")
    try:
        u=urlparse(raw); rd=parse_qs(u.query).get("rd"); host=urlparse(unquote(rd[0])).netloc if rd else u.netloc
    except Exception: host=""
    if "samsclub.com" in host: return "sams"
    if "walmart.com" in host: return "walmart"
    if "smithsfoodanddrug.com" in host: return "smiths"
    if "albertsons.com" in host: return "albertsons"
    if "safeway.com" in host: return "safeway"
    return None

PARSERS={"walmart":import_walmart.parse,"sams":import_sams.parse,"smiths":import_smiths.parse,"albertsons":import_albertsons.parse,"safeway":import_safeway.parse}

def main(src,dst):
    wb=load_workbook(src,read_only=True,data_only=True)
    observations=[]; rejected=[]; counts={}
    for ws in wb.worksheets:
        rows=ws.iter_rows(values_only=True)
        try: headers=[str(x) if x is not None else "" for x in next(rows)]
        except StopIteration: continue
        for n,values in enumerate(rows,2):
            row={headers[i]:values[i] for i in range(min(len(headers),len(values)))}
            retailer=retailer_for(row); parser=PARSERS.get(retailer)
            try: item=parser(row,Path(src).name,ws.title,n) if parser else None
            except Exception as e: item=None
            if item:
                observations.append(item); counts[item["retailer"]]=counts.get(item["retailer"],0)+1
            else: rejected.append({"sheet":ws.title,"row":n})
    Path(dst).write_text(json.dumps({"observations":observations,"audit":{"source":Path(src).name,"accepted":len(observations),"rejected":len(rejected),"stores":counts}},indent=2))
if __name__=="__main__": main(sys.argv[1],sys.argv[2])
